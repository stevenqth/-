"""Format existing scientific content by cloning the retained MathorCup template.
No model/data changes. Process appendix is deliberately excluded.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
import re,json,hashlib
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm,Pt

BASE=Path(__file__).resolve().parents[1]
TEMPLATE=BASE.parents[1]/'资料/Mathorcup论文标准模板(1).docx'
SOURCE=BASE/'output/literature_revision/论文正文_七篇文献优化版.md'
DEST=BASE/'output/literature_revision/RYH优化1.docx'


def parse():
    parts=[p.strip() for p in SOURCE.read_text(encoding='utf-8').split('\n\n') if p.strip()]
    sections={};name='preamble';sections[name]=[]
    for block in parts:
        if block.startswith('## '):
            name=block[3:];sections[name]=[]
        elif block.startswith('|') and sections[name] and sections[name][-1].startswith('|'):
            sections[name][-1]+='\n'+block
        else:sections[name].append(block)
    return sections


def main():
    original_hash=hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    ref=Document(TEMPLATE);doc=Document(TEMPLATE)
    samples={k:deepcopy(ref.paragraphs[i]._p.pPr) for k,i in [('title',2),('abstract',3),('body',5),('keyword',26),('h1',29),('h2',30),('h3',54),('reference',109)]}
    sections=parse();body=doc.element.body
    table_sequence=[1]
    for child in list(body):body.remove(child)
    def no_number(p):
        n=p._p.get_or_add_pPr().get_or_add_numPr();i=n.get_or_add_numId();i.val=0
    def para(text,role='body',style=None):
        p=doc.add_paragraph(style=style or ('Normal' if role in ['title','abstract','body','keyword'] else {'h1':'Heading 1','h2':'Heading 2','h3':'Heading 3','reference':'参考文献'}.get(role,'Normal')))
        if samples.get(role) is not None:
            if p._p.pPr is not None:p._p.remove(p._p.pPr)
            p._p.insert(0,deepcopy(samples[role]))
        p.add_run(text)
        if role in ['h1','h2','h3']:
            p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if role=='title':p.runs[0].font.size=Pt(14)
        if role=='keyword':p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        return p
    def heading(text,level=1):
        p=para(text,'h'+str(level))
        if level==1:p.paragraph_format.first_line_indent=Pt(0)
        return p
    def table(block):
        lines=[ln for ln in block.splitlines() if ln.strip().startswith('|')]
        rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
        rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c.strip()) for c in r)]
        n=len(rows[0]);t=doc.add_table(rows=len(rows),cols=n);t.style='三线表'
        # Template's 80% symbol-table width expands to usable page width for longer result tables.
        t.autofit=False; widths=[160/n]*n
        if n==2:widths=[34,126] if rows[0][0] in ['属性','维度'] else [105,55]
        if n==3:widths=[22,72,66] if rows[0][0] in ['目标','符号'] else ([25,55,80] if rows[0][0]=='来源' else [65,45,50])
        if n==4:widths=[55,25,25,55] if rows[0][0]=='字段' else [60,33,33,34]
        if n==5:widths=[50,25,28,28,29]
        for i,row in enumerate(rows):
            for j,value in enumerate(row):
                cell=t.cell(i,j);cell.width=Mm(widths[j]);p=cell.paragraphs[0]
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                indent=p._p.get_or_add_pPr().get_or_add_ind();indent.set(qn('w:firstLineChars'),'0');indent.set(qn('w:firstLine'),'0')
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or len(value)<15 else WD_ALIGN_PARAGRAPH.LEFT
                p.add_run(value)
                suppress=OxmlElement('w:suppressAutoHyphens');p._p.get_or_add_pPr().append(suppress)
            if i==0:
                repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            avoid=OxmlElement('w:cantSplit');t.rows[i]._tr.get_or_add_trPr().append(avoid)
        for c,w in zip(t.columns,widths):c.width=Mm(w)
    def blocks(items,sublevel=3,strip_subheads=False):
        for b in items:
            if b.startswith('### '):
                if not strip_subheads:heading(re.sub(r'^\d+\.\d+\s*','',b[4:]),sublevel)
            elif b.startswith('|'):table(b)
            elif b.startswith('!['):
                match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',b)
                if not match:raise ValueError(b)
                caption,path=match.groups();path=(SOURCE.parent/path).resolve()
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
                p.add_run().add_picture(str(path),width=Mm(145))
                cp=doc.add_paragraph(caption,style='图注');no_number(cp);cp.paragraph_format.keep_with_next=False
            elif re.match(r'^表\d+',b):
                table_sequence[0]+=1
                b=re.sub(r'^表\d+',f'表{table_sequence[0]}',b)
                p=doc.add_paragraph(b,style='表注');no_number(p);p.paragraph_format.keep_with_next=True
            else:para(b)
    def subsection(items,label):
        out=[];inside=False
        for b in items:
            if b.startswith('### '):
                if inside:break
                inside=label in b
            elif inside:out.append(b)
        return out
    def before_results(items):
        idx=next((i for i,b in enumerate(items) if b.startswith('|') or b.startswith('表') or b.startswith('![')),len(items))
        # Include a table's immediately preceding caption in the results group.
        return items[:idx],items[idx:]
    # Source-derived cover: keep table and two spacer paragraphs.
    body.append(deepcopy(ref.tables[0]._tbl))
    doc.tables[0].cell(1,1).text='A'
    body.append(deepcopy(ref.paragraphs[0]._p));body.append(deepcopy(ref.paragraphs[1]._p))
    para('啤酒感官预测与受约束风格迁移','title');para('摘 要','abstract')
    summary=sections['摘要']
    for b in summary:
        if b.startswith('关键词'):
            p=para('关键词：嵌套验证；条件重要性；稳定性选择；条件生成；Pareto决策','keyword')
        else:
            p=para(b)
            if b.startswith(('针对问题一','针对问题二','针对问题三')):
                # Template emphasizes each problem lead-in in the abstract.
                text=p.text;p.clear();p.add_run(text[:6]).bold=True;p.add_run(text[6:])
    body.append(deepcopy(ref.paragraphs[27]._p))
    # Keep original two-level/three-level TOC style via an updatable TOC field.
    p=doc.add_paragraph('目录');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.runs[0].font.name='黑体';p.runs[0].font.size=Pt(16)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
    for kind in ['begin','instr','separate','text','end']:
        run=OxmlElement('w:r')
        if kind=='instr':
            node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=' TOC \\o "1-3" \\h \\z \\u '
        elif kind=='text':node=OxmlElement('w:t');node.text='目录将在Word中更新'
        else:node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind)
        run.append(node);p._p.append(run)
    body.append(deepcopy(ref.paragraphs[28]._p))
    heading('问题重述');heading('问题背景',2)
    blocks(sections['一 问题与研究范围'])
    heading('问题提出',2)
    for text in ['问题一：利用化学与理化测量预测50项感官响应，并在评分锚点不足时提供明确标注的相对等级。',
                 '问题二：结合共表达网络筛选各感官属性的Top-3单项化学候选，并检验相关变量归因及名单稳定性。',
                 '问题三：在风格条件下生成化学画像，实现受约束跨风格迁移，并通过Pareto集合分析多目标取舍。']:para(text)
    heading('问题分析')
    for text in ['问题一的主要困难是样本量小、输入相关且量纲不统一。模型选择应隔离于最终误差评价，输出相对等级时不能把空白评分卡当作绝对评分标签。',
                 '问题二需要区分预测关联、名单稳定性和因果作用。相关化合物可能互相替代，因此除固定Top-3外，还需要条件置换与子抽样诊断。',
                 '问题三同时涉及生成分布、源样品改变和缺陷约束。第一问的代理误差会传递到迁移评价，故将尺度证据、约束可行性与真实品评效果分开讨论。']:para(text)
    heading('模型假设与符号说明');heading('模型基本假设',2)
    for text in ['同一beer编号代表同一款测量样品，化学与感官数据按该键对齐；风格标签一致性由输入数据核对。',
                 '训练与验证划分评价当前附件总体内的插值能力，不预设其能够外推到新酒厂、新批次或未见风格。',
                 '秩正态变换、近邻条件置换及风格收缩分布是统计建模假设，不赋予未知化学列确定的物理含义。',
                 '感官代理及Pareto候选用于筛选后续实验对象，不等同于已经完成真实加标、盲评或酿造验证。']:para(text)
    heading('符号说明',2)
    no_number(doc.add_paragraph('表1 符号说明',style='表注'))
    table('| 符号 | 含义 | 单位或范围 |\n| --- | --- | --- |\n| X | 化学与理化输入矩阵 | 250×231，附件分析尺度 |\n| Y | 感官响应矩阵 | 250×50，标准化值 |\n| s | 风格标签 | 22类 |\n| z | 条件VAE潜变量 | 8维，无量纲 |\n| pi | 子样本入选频率 | 0至1 |\n| f1 f2 f3 | 感官偏差、改动与缺陷指标 | 均取最小值 |')
    heading('数据预处理');heading('指标选取',2)
    data=sections['二 数据尺度与证据']
    blocks(subsection(data,'2.1'))
    heading('数据清洗',2);blocks(subsection(data,'2.2'));blocks(subsection(data,'2.3'))
    heading('模型建立与求解')
    heading('问题一模型建立与求解',2);heading('问题一求解思路',3)
    para('先在训练折内建立缺失与底值状态，再比较稳健线性和树模型候选；用内层选型、外层评价控制选择偏差，最后给出连续感官预测及相对等级。')
    one=sections['三 问题一 感官预测与验证隔离'];models,results=before_results(subsection(one,'3.1'))
    heading('问题一模型建立',3);blocks(models)
    heading('问题一模型求解与分析',3);blocks(results);blocks(subsection(one,'3.2'))
    heading('问题二模型建立与求解',2);heading('问题二求解思路',3)
    two=sections['四 问题二 相关化合物与稳定性'];blocks(subsection(two,'4.1'))
    model2,res2=before_results(subsection(two,'4.2'));stab=subsection(two,'4.3')
    heading('问题二模型建立',3);blocks(model2);blocks(stab[:1])
    heading('问题二模型求解与分析',3);blocks(res2);blocks(stab[1:])
    heading('问题三模型建立与求解',2);heading('问题三求解思路',3)
    para('对固定条件VAE的采样方式进行诊断，并用独立分层留出检验比较条件收缩高斯和经验基准；在潜空间中用NSGA-II搜索可行候选，再进行偏好后置决策。')
    three=sections['五 问题三 条件生成与分布检验'];multi=sections['六 多目标迁移与Pareto决策']
    heading('问题三模型建立',3);blocks(subsection(three,'5.1'))
    m3,r3=before_results(subsection(three,'5.2'));blocks(m3)
    m4,r4=before_results(subsection(three,'5.3'));blocks(m4)
    blocks(subsection(multi,'6.1'))
    heading('问题三模型求解与分析',3);blocks(r3);blocks(r4);blocks(subsection(multi,'6.2'))
    heading('模型分析检验');heading('灵敏度分析',2);blocks(subsection(multi,'6.3'))
    heading('误差分析',2);blocks(sections['七 外部数据对照与验证边界'])
    heading('模型评价与推广');heading('模型的优点',2)
    para('模型将专业评分、评论词频及化学尺度分别处理，使用重复嵌套验证评价完整选型程序，公开逐属性误差、条件重要性和子样本稳定性。约束迁移保留多目标取舍，使不同偏好可以对应不同候选，而不以单一总分掩盖差异。')
    heading('模型的不足',2)
    blocks(subsection(sections['八 局限与验证方案'],'8.2'))
    heading('模型的推广',2);blocks(subsection(sections['八 局限与验证方案'],'8.3'));blocks(sections['九 结论'])
    body.append(deepcopy(ref.paragraphs[96]._p))
    p=heading('参考文献');no_number(p)
    for b in sections['参考文献']:
        p=para(b,'reference');no_number(p);p.paragraph_format.first_line_indent=Pt(0)
    body.append(deepcopy(ref.sections[3]._sectPr))
    # Keep the first three section breaks and references section; no process-only appendix.
    doc.core_properties.title='RYH优化1'
    doc.core_properties.author=''
    settings=doc.settings.element
    update=settings.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');settings.append(update)
    update.set(qn('w:val'),'true')
    doc.save(DEST)
    # Restore opaque template parts byte-for-byte rather than normalizing their XML.
    with ZipFile(TEMPLATE) as z:original={n:z.read(n) for n in z.namelist()}
    with ZipFile(DEST) as z:new={n:z.read(n) for n in z.namelist()}
    preserve=[n for n in original if n in ['word/styles.xml','word/numbering.xml','word/theme/theme1.xml'] or n.startswith('word/header') or n.startswith('word/footer')]
    for n in preserve:new[n]=original[n]
    with ZipFile(DEST,'w',ZIP_DEFLATED) as z:
        for n,data in new.items():z.writestr(n,data)
    assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()==original_hash
    checks={'template_sha256':original_hash,'source':str(SOURCE),'output':str(DEST),
            'sections':len(Document(DEST).sections),'keyword_alignment':'left',
            'appendix_process_removed':True,'preserved_template_parts':preserve}
    (BASE/'results/literature_revision/RYH优化1_格式核对.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    print(DEST)


if __name__=='__main__':main()
