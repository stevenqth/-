# -*- coding: utf-8 -*-
"""08_final_check.py — 论文与提交包终检（T016 收尾）。

检查项：
    1) 论文 docx：可打开、段落/表格/图片计数、章节齐全、摘要含“针对问题一/二/三”；
    2) 论文 PDF：存在、%PDF 签名、非空；
    3) 分章 Markdown 引用的图片全部存在；
    4) 成果/results/*.json 全部可解析；
    5) 提交包完整性：文件数、提交清单 CSV 的 SHA256 与实际一致、zip 存在；
    6) 占位符检查：MCB250XXXX（预期存在，待官方编号）。
输出：成果/02_论文/终检清单.md

运行（需 python-docx；使用工作区依赖运行时）：
    <python-with-docx> 成果/code/08_final_check.py
"""
import os
import re
import csv
import json
import hashlib

from docx import Document

ROOT = r"D:\github mama"
PAPER = os.path.join(ROOT, "成果", "02_论文")
FIG = os.path.join(ROOT, "成果", "figures")
RES = os.path.join(ROOT, "成果", "results")
PKG = os.path.join(ROOT, "成果", "提交包")
ZIP = os.path.join(ROOT, "成果", "MathorCup2026_赛道A_参赛包_占位.zip")
OUT = os.path.join(PAPER, "终检清单.md")

rows = []


def check(name, ok, detail=""):
    rows.append((name, "PASS" if ok else "FAIL", detail))
    print(("PASS " if ok else "FAIL ") + name + (" | " + detail if detail else ""))


docx_path = os.path.join(PAPER, "00_论文全文.docx")
pdf_path = os.path.join(PAPER, "00_论文全文.pdf")
check("论文 docx 存在", os.path.exists(docx_path),
      f"{os.path.getsize(docx_path)/1024/1024:.2f} MB" if os.path.exists(docx_path) else "缺失")
check("论文 PDF 存在", os.path.exists(pdf_path),
      f"{os.path.getsize(pdf_path)/1024/1024:.2f} MB" if os.path.exists(pdf_path) else "缺失")

if os.path.exists(pdf_path):
    with open(pdf_path, "rb") as f:
        sig = f.read(5)
    check("PDF 签名 %PDF-", sig == b"%PDF-", str(sig))

doc_text = ""
if os.path.exists(docx_path):
    d = Document(docx_path)
    doc_text = "\n".join(p.text for p in d.paragraphs)
    check("docx 段落数 ≥ 300", len(d.paragraphs) >= 300, f"{len(d.paragraphs)} 段")
    check("docx 表格数 ≥ 20", len(d.tables) >= 20, f"{len(d.tables)} 表")
    check("docx 图片数 ≥ 15", len(d.inline_shapes) >= 15, f"{len(d.inline_shapes)} 图")
    chapters = ["摘 要", "问题重述", "问题分析", "模型假设与符号说明", "数据预处理",
                "问题一", "问题二", "问题三", "模型分析与检验", "模型评价与推广",
                "参考文献", "附录"]
    missing = [c for c in chapters if c not in doc_text]
    check("论文章节齐全", not missing, "缺失:" + ",".join(missing) if missing else "12 个章节均存在")
    check("摘要含“针对问题一/二/三”",
          all(k in doc_text for k in ["针对问题一", "针对问题二", "针对问题三"]),
          "摘要四段式")
    check("占位符 MCB250XXXX 存在（待官方编号）", "MCB250XXXX" in doc_text,
          "预期占位，官方编号到位后替换")

# 图片引用检查
refs = set()
for fn in os.listdir(PAPER):
    if fn.endswith(".md") and re.match(r"^(0\d|1\d)_", fn):
        with open(os.path.join(PAPER, fn), encoding="utf-8") as f:
            for m in re.finditer(r"!\[[^\]]*\]\((\.\./figures/[^)]+)\)", f.read()):
                refs.add(os.path.normpath(os.path.join(PAPER, m.group(1))))
missing_figs = [p for p in refs if not os.path.exists(p)]
check("分章 Markdown 图片引用均存在", not missing_figs,
      f"{len(refs)} 个引用" + ("；缺失:" + ";".join(missing_figs) if missing_figs else ""))

# 结果 JSON
json_files = [f for f in os.listdir(RES) if f.endswith(".json")]
bad = []
for f in json_files:
    try:
        with open(os.path.join(RES, f), encoding="utf-8") as fh:
            json.load(fh)
    except Exception as e:
        bad.append(f + ":" + str(e)[:40])
check("results/*.json 全部可解析", not bad, f"{len(json_files)} 个 JSON" + ("；失败:" + ";".join(bad) if bad else ""))

# 提交包
pkg_files = []
for dirpath, _, files in os.walk(PKG):
    for f in files:
        if not f.endswith(".zip"):
            pkg_files.append(os.path.join(dirpath, f))
check("提交包文件数 ≥ 15", len(pkg_files) >= 15, f"{len(pkg_files)} 个文件")
check("提交包 zip 存在", os.path.exists(ZIP),
      f"{os.path.getsize(ZIP)/1024/1024:.2f} MB" if os.path.exists(ZIP) else "缺失")

checklist = os.path.join(PKG, "提交清单.csv")
if os.path.exists(checklist):
    with open(checklist, encoding="utf-8-sig") as f:
        rows_csv = list(csv.DictReader(f))
    mism = []
    for r in rows_csv:
        p = os.path.join(PKG, r["path"])
        if not os.path.exists(p):
            mism.append(r["path"] + ":缺失")
            continue
        h = hashlib.sha256(open(p, "rb").read()).hexdigest().upper()
        if h != r["sha256"].upper():
            mism.append(r["path"] + ":hash不一致")
    check("提交清单 SHA256 与实际一致", not mism, f"{len(rows_csv)} 条" + ("；问题:" + ";".join(mism) if mism else ""))

with open(OUT, "w", encoding="utf-8") as f:
    f.write("# 论文与提交包终检清单\n\n")
    f.write("> 生成脚本：`成果/code/08_final_check.py`；生成时间：2026-10-28（Asia/Shanghai）。\n\n")
    f.write("| 检查项 | 结果 | 说明 |\n| --- | --- | --- |\n")
    for n, s, d in rows:
        f.write(f"| {n} | {s} | {d} |\n")
    fails = [r for r in rows if r[1] == "FAIL"]
    f.write(f"\n**结论**：共 {len(rows)} 项检查，PASS {len(rows)-len(fails)} 项，FAIL {len(fails)} 项。")
    if fails:
        f.write("\n\n需处理：" + "；".join(f"`{r[0]}`（{r[2]}）" for r in fails))
    else:
        f.write(" 论文与提交包结构完整，可进入提交环节（封面队伍编号与文件命名待官方信息替换）。")
print("saved:", OUT)
raise SystemExit(0 if not fails else 1)
