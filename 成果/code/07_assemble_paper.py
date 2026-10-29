# -*- coding: utf-8 -*-
"""07_assemble_paper.py — 按官方模板结构把分章 Markdown 汇编为 docx（T016）。

使用 python-docx（需用带 python-docx 的 Python 解释器运行，例如工作区依赖运行时）：
    <python-with-docx> 成果/code/07_assemble_paper.py
输出：
    成果/02_论文/00_论文全文.docx

解析规则（够用即可）：
    #/##/###/#### 标题；| ... | 表格（含分隔行）；![alt](path) 图片；``` 代码块；
    -/* 项目符号；1. 有序列表；其余为段落（支持 **粗体**、`代码`、[文字](链接) 的内联清洗）。
"""
import os
import re
import sys
import datetime

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"D:\github mama"
PAPER = os.path.join(ROOT, "成果", "02_论文")
FIG = os.path.join(ROOT, "成果", "figures")
OUT = os.path.join(PAPER, "00_论文全文.docx")

TITLE = "基于多输出树集成与条件生成模型的啤酒风味预测、解析与设计"
TEAM_NO = "MCB250XXXX"
TRACK = "赛道 A"

ORDER = [
    "00_摘要与关键词.md",
    "01_问题重述与问题分析.md",
    "02_模型假设与符号说明.md",
    "03_数据预处理.md",
    "03_问题一_模型与结果.md",
    "04_问题二_模型与结果.md",
    "05_问题三_模型与结果.md",
    "06_模型分析与检验.md",
    "07_模型评价与推广.md",
    "08_参考文献.md",
    "09_附录.md",
]

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
LATIN_FONT = "Times New Roman"
CODE_FONT = "Consolas"


def set_run_font(run, name=BODY_FONT, size=12, bold=None, color=None):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), LATIN_FONT)
    rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    rFonts.set(qn("w:eastAsia"), name)


def set_style_font(doc, style_name, name, size, bold=False):
    try:
        st = doc.styles[style_name]
    except KeyError:
        return
    st.font.name = LATIN_FONT
    st.font.size = Pt(size)
    st.font.bold = bold
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), LATIN_FONT)
    rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    rFonts.set(qn("w:eastAsia"), name)


def clean_inline(text):
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    return text.strip()


def add_paragraph(doc, text, style=None, size=12, indent_first=True, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent_first and style is None:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(clean_inline(text))
    set_run_font(run, BODY_FONT, size)
    return p


def add_heading(doc, text, level):
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}.get(level, "Heading 4")
    p = doc.add_paragraph(style=style)
    sizes = {1: 16, 2: 14, 3: 13, 4: 12}
    run = p.add_run(clean_inline(text))
    set_run_font(run, HEAD_FONT, sizes.get(level, 12), bold=True, color=RGBColor(0, 0, 0))
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_table(doc, rows):
    n_cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n_cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = t.cell(i, j)
            cell.text = ""
            val = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(clean_inline(val))
            set_run_font(run, BODY_FONT, 10.5, bold=(i == 0))
    doc.add_paragraph()  # 表后空行


def add_image(doc, path, caption, width_cm=14.5):
    if not os.path.exists(path):
        add_paragraph(doc, f"[缺图：{os.path.basename(path)}]", indent_first=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, BODY_FONT, 10.5)
    cap.paragraph_format.space_after = Pt(6)


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("\n".join(lines))
    run.font.name = CODE_FONT
    run.font.size = Pt(8)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), CODE_FONT)
    rFonts.set(qn("w:hAnsi"), CODE_FONT)
    rFonts.set(qn("w:eastAsia"), BODY_FONT)


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def parse_markdown(doc, md_path, fig_counter, prefix=None):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    i, para_buf = 0, []
    h1_used = [False]

    def flush():
        nonlocal para_buf
        if para_buf:
            text = " ".join(x.strip() for x in para_buf).strip()
            if text:
                add_paragraph(doc, text)
            para_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush(); i += 1; continue
        # 代码块
        if line.startswith("```"):
            flush()
            code, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            add_code(doc, code)
            continue
        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            flush()
            heading_text = m.group(2)
            if len(m.group(1)) == 1 and prefix and not h1_used[0]:
                heading_text = prefix + heading_text
                h1_used[0] = True
            add_heading(doc, heading_text, len(m.group(1)))
            i += 1; continue
        # 表格
        if line.startswith("|"):
            flush()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip()); i += 1
            rows = []
            for r in table_lines:
                if re.match(r"^\|[\s:\-|]+\|$", r):
                    continue
                cells = [c.strip() for c in r.strip().strip("|").split("|")]
                rows.append(cells)
            if rows:
                add_table(doc, rows)
            continue
        # 图片
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", line)
        if m:
            flush()
            fig_counter[0] += 1
            alt, rel = m.group(1), m.group(2)
            path = os.path.normpath(os.path.join(PAPER, rel))
            add_image(doc, path, f"图 {fig_counter[0]}  {clean_inline(alt)}")
            i += 1; continue
        # 分隔线
        if re.match(r"^-{3,}$", line) or re.match(r"^\*{3,}$", line):
            flush(); i += 1; continue
        # 引用块（章节顶部的元信息等）→ 小字灰色普通段落
        if line.startswith(">"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(clean_inline(line.lstrip(">").strip()))
            set_run_font(r, BODY_FONT, 10.5, color=RGBColor(0x66, 0x66, 0x66))
            i += 1; continue
        # 列表
        if re.match(r"^[-*]\s+", line):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(clean_inline(re.sub(r"^[-*]\s+", "", line)))
            set_run_font(run, BODY_FONT, 12)
            i += 1; continue
        if re.match(r"^\d+\.\s+", line):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(0)
            run = p.add_run(clean_inline(line))
            set_run_font(run, BODY_FONT, 12)
            i += 1; continue
        # 引用/普通段落
        para_buf.append(line)
        i += 1
    flush()


def build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026 年 MathorCup 数学应用挑战赛——大数据竞赛")
    set_run_font(r, HEAD_FONT, 15, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("（赛道 A：啤酒酿造化合物分析与啤酒风味预测）")
    set_run_font(r, HEAD_FONT, 13, bold=True)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    set_run_font(r, HEAD_FONT, 20, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    for label, value in [("队 伍 编 号", TEAM_NO), ("参 赛 赛 道", TRACK)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}：{value}")
        set_run_font(r, BODY_FONT, 14)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026 年 10 月")
    set_run_font(r, BODY_FONT, 13)
    doc.add_page_break()


def build_toc(doc):
    add_heading(doc, "目  录", 1)
    entries = [
        ("摘 要", ""), ("一、问题重述", ""), ("二、问题分析", ""),
        ("三、模型假设与符号说明", ""), ("四、数据预处理", ""),
        ("五、问题一：数据缺失条件下的感官评分", ""),
        ("六、问题二：化合物–风味关键因素识别与稳定性", ""),
        ("七、问题三：风格条件生成与跨风格迁移", ""),
        ("八、模型分析与检验：灵敏度与稳健性", ""),
        ("九、模型评价与推广", ""), ("十、参考文献", ""), ("十一、附录", ""),
    ]
    for name, _ in entries:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(name)
        set_run_font(r, BODY_FONT, 12)
    doc.add_page_break()


def main():
    doc = Document()
    # 页面设置 A4
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    # 样式
    set_style_font(doc, "Normal", BODY_FONT, 12)
    for lvl, size in [(1, 16), (2, 14), (3, 13), (4, 12)]:
        try:
            set_style_font(doc, f"Heading {lvl}", HEAD_FONT, size, bold=True)
        except Exception:
            pass
    add_page_number_footer(doc)

    build_cover(doc)
    build_toc(doc)

    fig_counter = [0]
    prefixes = {
        "03_问题一_模型与结果.md": "五、",
        "04_问题二_模型与结果.md": "六、",
        "05_问题三_模型与结果.md": "七、",
        "06_模型分析与检验.md": "八、",
    }
    for idx, fname in enumerate(ORDER):
        path = os.path.join(PAPER, fname)
        if not os.path.exists(path):
            print("MISSING:", fname)
            continue
        if idx > 0:
            doc.add_page_break()
        parse_markdown(doc, path, fig_counter, prefixes.get(fname))
        print(f"parsed {fname}", flush=True)

    doc.save(OUT)
    print("saved:", OUT, flush=True)
    # 校验
    d2 = Document(OUT)
    print("paragraphs:", len(d2.paragraphs), "tables:", len(d2.tables),
          "inline_shapes:", len(d2.inline_shapes), "size_kb:",
          round(os.path.getsize(OUT) / 1024, 1), flush=True)
    print("figures inserted:", fig_counter[0], flush=True)


if __name__ == "__main__":
    main()
